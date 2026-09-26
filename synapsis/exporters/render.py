"""
Shared Markdown renderer for every export surface (review L6-02, 2026-09-26).

Before this module the HTML/PDF exports escaped the assistant's Markdown and
swapped ``\\n`` for ``<br>``, and the DOCX export wrote each paragraph as plain
text: readers saw pipe tables, ``**bold**``, non-clickable ``[R1003](…)``
citations and raw ``<chart>{json}</chart>`` specs. Now:

* :func:`prepare_assistant_text` — the text pipeline applied to every assistant
  message before rendering: control characters stripped (DOCX XML safety,
  L6-12), result codes linked to their public PRMS report (Lane G's
  ``linkify_result_codes``, idempotent, so chats saved before the linkifier
  also export with links) and internal server paths reduced to the file name.
* :func:`markdown_to_html` — ``markdown-it-py`` (CommonMark + GFM tables and
  strikethrough) with **raw HTML disabled** (model text can never inject
  markup), images disabled (a PDF render must never fetch remote content),
  bare ``https://`` URLs auto-linked, and only ``http``/``https``
  links kept as links (anything else — relative paths, ``mailto:``, ``file:``,
  ``javascript:`` — is rendered as its plain text).
* :func:`markdown_into_docx` — the same token stream written into a
  python-docx document: headings, paragraphs with bold/italic/strike/code
  runs, real hyperlinks, bullet and numbered lists, tables, code blocks and
  quotes.
* ``<chart>{spec}</chart>`` blocks become a captioned data table in every
  format (:func:`split_chart_blocks`, :func:`chart_to_table`).

The zero-draft watermark, disclaimer box, per-page diagonal mark and snapshot
line are NOT produced here — they stay in :mod:`synapsis.exporters.watermark`,
unchanged in wording.
"""

from __future__ import annotations

import html as _html
import json
import re
from typing import Any, Iterable, Optional

from markdown_it import MarkdownIt
from markdown_it.token import Token

# ---------------------------------------------------------------------------
# Text preparation
# ---------------------------------------------------------------------------

#: XML 1.0 forbids most C0 control characters; python-docx raises on them
#: (review L6-12) and they are noise in every other format too.
_CONTROL = re.compile(r"[\x00-\x08\x0b\x0c\x0e-\x1f\x7f]")

#: ANSI escape sequences (tool output pasted into answers).
_ANSI = re.compile(r"\x1b\[[0-9;?]*[ -/]*[@-~]")

#: Files in a user's own area: ``…/(uploads|outputs)/u/<24 hex>/[<32 hex>/]name``.
_OWNER_AREA_PATH = re.compile(
    r"(?<![\w/.:\-])(?:file://)?(?:/[\w.\-]+)*/(?:uploads|outputs)/u/[0-9a-f]{24}/(?:[0-9a-f]{32}/)?"
    r"(?P<name>[^\s/)\]`'\"<>|]+)"
)

#: Any other absolute server path under the container workspace or a home dir.
_SERVER_PATH = re.compile(
    r"(?<![\w/.:\-])(?:file://)?/(?:workspace|app|tmp|home/[\w.\-]+|Users/[\w.\-]+)(?:/[\w.\-]+)*/"
    r"(?P<name>[\w.\-]+\.[A-Za-z0-9]{1,6})(?![\w/])"
)

_FILE_SCHEME = re.compile(r"file://", re.IGNORECASE)


def clean_text(text: Any) -> str:
    """Strip ANSI sequences and XML-invalid control characters."""
    s = "" if text is None else str(text)
    return _CONTROL.sub("", _ANSI.sub("", s))


def scrub_internal_paths(text: str) -> str:
    """Reduce server file paths to the bare file name (exports leave the app).

    ``/workspace/outputs/u/<owner>/<token>/Report.docx`` → ``Report.docx``.
    The download link only works inside the app, and the path exposes the
    server layout; the reader still sees which file the answer referred to.
    """
    text = _OWNER_AREA_PATH.sub(lambda m: m.group("name"), text)
    text = _SERVER_PATH.sub(lambda m: m.group("name"), text)
    return _FILE_SCHEME.sub("", text)


def linkify(text: str) -> str:
    """Link result codes to their public PRMS report (Lane G). Never raises."""
    try:
        from synapsis.tools.result_code_citation import linkify_result_codes

        return linkify_result_codes(text)
    except Exception:  # noqa: BLE001 — an export must never fail on linking
        return text


def prepare_assistant_text(text: Any) -> str:
    """The full pipeline for an assistant message before any renderer."""
    return scrub_internal_paths(linkify(clean_text(text)))


def prepare_user_text(text: Any) -> str:
    """User messages: cleaned and scrubbed, but never rewritten (no linking)."""
    return scrub_internal_paths(clean_text(text))


# ---------------------------------------------------------------------------
# <chart> blocks
# ---------------------------------------------------------------------------

_CHART_BLOCK = re.compile(r"<chart>\s*(?P<body>.*?)\s*</chart>", re.DOTALL | re.IGNORECASE)
_CHART_FENCE = re.compile(r"```(?:json|chart)?\s*\n(?P<body>\{.*?\})\s*\n```", re.DOTALL)


def _parse_chart(body: str) -> Optional[dict]:
    try:
        spec = json.loads(body)
    except (ValueError, TypeError):
        return None
    if not isinstance(spec, dict) or not isinstance(spec.get("data"), list):
        return None
    return spec


def split_chart_blocks(text: str) -> list[tuple[str, Any]]:
    """Split *text* into ``("md", str)`` and ``("chart", spec|None)`` segments.

    Explicit ``<chart>`` blocks always become a chart segment (``None`` when
    the JSON is unreadable). A fenced JSON block counts only when it carries
    ``chartType`` + ``data`` (the frontend renders those as charts too).
    """
    segments: list[tuple[str, Any]] = []
    pos = 0
    matches = sorted(
        [m for m in _CHART_BLOCK.finditer(text)]
        + [m for m in _CHART_FENCE.finditer(text)
           if (lambda s: s is not None and "chartType" in s)(_parse_chart(m.group("body")))],
        key=lambda m: m.start(),
    )
    for m in matches:
        if m.start() < pos:
            continue  # overlapping (a fence inside a <chart> block)
        if m.start() > pos:
            segments.append(("md", text[pos:m.start()]))
        segments.append(("chart", _parse_chart(m.group("body"))))
        pos = m.end()
    if pos < len(text):
        segments.append(("md", text[pos:]))
    return [s for s in segments if s[0] == "chart" or s[1].strip()]


def format_value(value: Any) -> str:
    """Chart cell → text: ``None`` = "n/a" (a gap, never 0); integers with separators."""
    if value is None:
        return "n/a"
    if isinstance(value, bool):
        return str(value)
    if isinstance(value, int):
        return f"{value:,}"
    if isinstance(value, float):
        if value != value:  # NaN
            return "n/a"
        if value.is_integer():
            return f"{int(value):,}"
        return f"{value:,.2f}".rstrip("0").rstrip(".")
    return clean_text(value)


_CHART_KIND = {
    "bar": "bar chart", "multiBar": "grouped bar chart", "line": "line chart",
    "area": "area chart", "stackedArea": "stacked area chart", "pie": "pie chart",
    "scatter": "scatter chart", "horizontalBar": "bar chart", "stackedBar": "stacked bar chart",
}


def chart_to_table(spec: Optional[dict]) -> tuple[str, list[str], list[list[str]]]:
    """Return ``(caption, columns, rows)`` for a chart spec (all strings).

    When the spec is unreadable the caption says so and the table is empty.
    """
    if not spec:
        return ("Chart omitted from this export: its data could not be read.", [], [])
    title = clean_text(spec.get("title") or "Chart").strip() or "Chart"
    kind = _CHART_KIND.get(str(spec.get("chartType") or ""), "chart")
    data = [d for d in spec.get("data") or [] if isinstance(d, dict)]
    x_key = spec.get("xAxisKey")
    series = [s for s in (spec.get("series") or []) if isinstance(s, dict) and s.get("key")]
    if not x_key and data:
        x_key = next((k for k, v in data[0].items() if isinstance(v, str)), None)
    if not series and data:
        series = [{"key": k} for k, v in data[0].items() if k != x_key and isinstance(v, (int, float))]
    columns: list[str] = []
    keys: list[str] = []
    if x_key:
        columns.append(str(x_key).replace("_", " ").strip().capitalize() or "Category")
        keys.append(str(x_key))
    for s in series:
        columns.append(clean_text(s.get("label") or str(s["key"]).replace("_", " ").title()))
        keys.append(str(s["key"]))
    rows = [[format_value(d.get(k)) if k != x_key else clean_text(d.get(k, "")) for k in keys] for d in data]
    caption = f"{title} ({kind}; shown as a data table in this export)"
    desc = clean_text(spec.get("description") or "").strip()
    if desc:
        caption += f" — {desc}"
    return caption, columns, rows


def chart_to_markdown(spec: Optional[dict]) -> str:
    """A chart as a Markdown caption + pipe table (for the .md export)."""
    caption, columns, rows = chart_to_table(spec)
    if not columns:
        return f"*{caption}*"

    def cell(v: str) -> str:
        return v.replace("|", "\\|").replace("\n", " ")

    lines = [f"*{caption}*", "", "| " + " | ".join(cell(c) for c in columns) + " |",
             "|" + "|".join(" --- " for _ in columns) + "|"]
    lines += ["| " + " | ".join(cell(v) for v in r) + " |" for r in rows]
    return "\n".join(lines)


def replace_charts_with_markdown_tables(text: str) -> str:
    """Markdown export: every ``<chart>`` block becomes a captioned pipe table."""
    out = []
    for kind, value in split_chart_blocks(text):
        out.append(value if kind == "md" else "\n\n" + chart_to_markdown(value) + "\n\n")
    return "".join(out)


# ---------------------------------------------------------------------------
# markdown-it configuration
# ---------------------------------------------------------------------------

#: Schemes kept as clickable links in exports (no mailto: exports carry no
#: contact addresses — Jose, 2026-08-09).
SAFE_LINK_SCHEMES = ("http://", "https://")

_BARE_URL = re.compile(r"https?://[^\s<>\"'`]+[^\s<>\"'`.,;:!?)\]]")


def is_safe_href(href: Optional[str]) -> bool:
    return bool(href) and href.strip().lower().startswith(SAFE_LINK_SCHEMES)


def _autolink_bare_urls(state) -> None:
    """Core rule: turn bare ``https://…`` in text tokens into link tokens.

    (``linkify-it-py`` would add a dependency for this one job.) Text inside
    an existing link is left alone.
    """
    for block in state.tokens:
        if block.type != "inline" or not block.children:
            continue
        new_children: list[Token] = []
        depth = 0
        for tok in block.children:
            if tok.type == "link_open":
                depth += 1
            elif tok.type == "link_close":
                depth = max(0, depth - 1)
            if tok.type != "text" or depth or "://" not in tok.content:
                new_children.append(tok)
                continue
            pos = 0
            for m in _BARE_URL.finditer(tok.content):
                if m.start() > pos:
                    t = Token("text", "", 0)
                    t.content = tok.content[pos:m.start()]
                    new_children.append(t)
                url = m.group(0)
                lo = Token("link_open", "a", 1)
                lo.attrs = {"href": url}
                lo.markup = "linkify"
                new_children.append(lo)
                t = Token("text", "", 0)
                t.content = url
                new_children.append(t)
                new_children.append(Token("link_close", "a", -1))
                pos = m.end()
            if pos == 0:
                new_children.append(tok)
                continue
            if pos < len(tok.content):
                t = Token("text", "", 0)
                t.content = tok.content[pos:]
                new_children.append(t)
        block.children = new_children


def _drop_unsafe_links(state) -> None:
    """Core rule: links whose target is not http(s) keep only their text."""
    for block in state.tokens:
        if block.type != "inline" or not block.children:
            continue
        stack: list[bool] = []
        kept: list[Token] = []
        for tok in block.children:
            if tok.type == "link_open":
                safe = is_safe_href(tok.attrGet("href"))
                stack.append(safe)
                if safe:
                    kept.append(tok)
            elif tok.type == "link_close":
                safe = stack.pop() if stack else True
                if safe:
                    kept.append(tok)
            else:
                kept.append(tok)
        block.children = kept


def _build_md() -> MarkdownIt:
    md = MarkdownIt("commonmark", {"html": False, "linkify": False, "typographer": False, "breaks": False})
    md.enable(["table", "strikethrough"])
    md.disable(["image"])  # never fetch remote content while rendering a PDF
    md.core.ruler.push("ia_autolink_bare_urls", _autolink_bare_urls)
    md.core.ruler.push("ia_drop_unsafe_links", _drop_unsafe_links)

    def link_open(self, tokens, idx, options, env):
        tokens[idx].attrSet("target", "_blank")
        tokens[idx].attrSet("rel", "noopener noreferrer")
        return self.renderToken(tokens, idx, options, env)

    md.add_render_rule("link_open", link_open)
    return md


_MD = _build_md()


def parse_markdown(text: str) -> list[Token]:
    return _MD.parse(text or "")


def markdown_to_html(text: str) -> str:
    """Render Markdown (no raw HTML) to an HTML fragment."""
    return _MD.render(text or "")


def chart_to_html(spec: Optional[dict]) -> str:
    caption, columns, rows = chart_to_table(spec)
    cap = f"<figcaption>{_html.escape(caption)}</figcaption>"
    if not columns:
        return f'<figure class="chart-table">{cap}</figure>'
    head = "".join(f"<th>{_html.escape(c)}</th>" for c in columns)
    body = "".join(
        "<tr>" + "".join(f"<td>{_html.escape(v)}</td>" for v in r) + "</tr>" for r in rows
    )
    return (
        f'<figure class="chart-table">{cap}'
        f"<table><thead><tr>{head}</tr></thead><tbody>{body}</tbody></table></figure>"
    )


def render_message_html(text: str) -> str:
    """A prepared message (Markdown + chart blocks) → HTML fragment."""
    parts = []
    for kind, value in split_chart_blocks(text):
        parts.append(markdown_to_html(value) if kind == "md" else chart_to_html(value))
    return "\n".join(parts)


#: CSS for rendered message bodies (appended to the HTML export's styles).
RENDERED_CSS = """\
  .md p { margin: 0.4rem 0; }
  .md ul, .md ol { margin: 0.4rem 0 0.4rem 1.4rem; padding: 0; }
  .md h1, .md h2, .md h3, .md h4 { margin: 0.9rem 0 0.4rem; line-height: 1.3; }
  .md h1 { font-size: 1.35rem; } .md h2 { font-size: 1.2rem; } .md h3 { font-size: 1.05rem; }
  .md a { color: #1565c0; text-decoration: underline; word-break: break-word; }
  .md table { font-size: 0.85rem; page-break-inside: auto; }
  .md tr { page-break-inside: avoid; }
  figure.chart-table { margin: 1rem 0; }
  figure.chart-table figcaption { font-weight: 600; font-size: 0.85rem; color: #2e7d32; margin-bottom: 0.3rem; }"""


# ---------------------------------------------------------------------------
# DOCX writer
# ---------------------------------------------------------------------------

_MONO = "Consolas"


def _add_hyperlink(paragraph, url: str, text: str, *, bold=False, italic=False, strike=False) -> None:
    """Append a real, clickable external hyperlink run to *paragraph*."""
    from docx.opc.constants import RELATIONSHIP_TYPE as RT
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn

    r_id = paragraph.part.relate_to(url, RT.HYPERLINK, is_external=True)
    link = OxmlElement("w:hyperlink")
    link.set(qn("r:id"), r_id)
    run = OxmlElement("w:r")
    rpr = OxmlElement("w:rPr")
    for tag, on in (("w:b", bold), ("w:i", italic), ("w:strike", strike)):
        if on:
            rpr.append(OxmlElement(tag))
    color = OxmlElement("w:color")
    color.set(qn("w:val"), "1565C0")
    rpr.append(color)
    underline = OxmlElement("w:u")
    underline.set(qn("w:val"), "single")
    rpr.append(underline)
    run.append(rpr)
    t = OxmlElement("w:t")
    t.set(qn("xml:space"), "preserve")
    t.text = text
    run.append(t)
    link.append(run)
    paragraph._p.append(link)


def _write_inline(paragraph, children: Iterable[Token], base_size=None) -> None:
    """Write markdown-it inline tokens as runs (bold/italic/strike/code/links)."""
    bold = italic = strike = False
    href: Optional[str] = None
    link_text: list[str] = []

    def emit(text: str, code: bool = False) -> None:
        if not text:
            return
        if href is not None:
            link_text.append(text)
            return
        run = paragraph.add_run(text)
        run.bold = bold or None
        run.italic = italic or None
        if strike:
            run.font.strike = True
        if code:
            run.font.name = _MONO
        if base_size:
            run.font.size = base_size

    for tok in children:
        t = tok.type
        if t == "text":
            emit(tok.content)
        elif t == "code_inline":
            emit(tok.content, code=True)
        elif t == "softbreak":
            emit(" ")
        elif t == "hardbreak":
            if href is None:
                paragraph.add_run().add_break()
        elif t == "strong_open":
            bold = True
        elif t == "strong_close":
            bold = False
        elif t == "em_open":
            italic = True
        elif t == "em_close":
            italic = False
        elif t == "s_open":
            strike = True
        elif t == "s_close":
            strike = False
        elif t == "link_open":
            href = tok.attrGet("href")
            link_text = []
        elif t == "link_close":
            if href is not None:
                text = "".join(link_text) or href
                try:
                    _add_hyperlink(paragraph, href, text, bold=bold, italic=italic, strike=strike)
                except Exception:  # noqa: BLE001 — never lose the text
                    paragraph.add_run(f"{text} ({href})")
            href = None
            link_text = []
    if href is not None and link_text:  # unbalanced (should not happen)
        paragraph.add_run("".join(link_text))


def _style_or_default(doc, name: str):
    try:
        return doc.styles[name]
    except KeyError:
        return None


def _add_docx_table(container, header: list, rows: list, *, header_tokens=None, row_tokens=None) -> None:
    """Add a grid table. Cells are plain strings, or inline-token lists."""
    from docx.shared import Pt

    ncols = max([len(header)] + [len(r) for r in rows] + [1])
    table = container.add_table(rows=1, cols=ncols)
    try:
        table.style = "Table Grid"
    except Exception:  # noqa: BLE001 — style missing in a custom template
        pass
    size = Pt(9)

    def fill(cell, value, is_header=False):
        para = cell.paragraphs[0]
        if isinstance(value, list):
            _write_inline(para, value, base_size=size)
            if is_header:
                for run in para.runs:
                    run.bold = True
        else:
            run = para.add_run(clean_text(value))
            run.bold = is_header or None
            run.font.size = size

    for i in range(ncols):
        fill(table.rows[0].cells[i], (header_tokens or header)[i] if i < len(header_tokens or header) else "", True)
    for r_i, r in enumerate(rows):
        cells = table.add_row().cells
        values = row_tokens[r_i] if row_tokens else r
        for i in range(ncols):
            fill(cells[i], values[i] if i < len(values) else "")


def _add_caption(container, caption: str) -> None:
    from docx.shared import Pt, RGBColor

    p = container.add_paragraph()
    run = p.add_run(caption)
    run.bold = True
    run.font.size = Pt(9)
    run.font.color.rgb = RGBColor(0x2E, 0x7D, 0x32)


def chart_into_docx(container, spec: Optional[dict]) -> None:
    caption, columns, rows = chart_to_table(spec)
    _add_caption(container, caption)
    if columns:
        _add_docx_table(container, columns, rows)


def _list_style(doc, ordered: bool, level: int) -> Optional[str]:
    base = "List Number" if ordered else "List Bullet"
    name = base if level <= 1 else f"{base} {min(level, 3)}"
    if doc is not None and _style_or_default(doc, name) is None:
        return base if _style_or_default(doc, base) is not None else None
    return name


def markdown_into_docx(doc, text: str, *, heading_offset: int = 0, base_size=None) -> None:
    """Write Markdown (plus ``<chart>`` blocks) into a python-docx ``Document``.

    Args:
        doc:            the ``Document`` (body container).
        text:           prepared Markdown.
        heading_offset: added to Markdown heading levels (a ``#`` inside an
                        exported chat message sits under the "Assistant"
                        level-2 heading, so the exporter passes 2).
        base_size:      optional ``Pt`` size for body runs.
    """
    for kind, value in split_chart_blocks(text or ""):
        if kind == "chart":
            chart_into_docx(doc, value)
        else:
            _tokens_into_docx(doc, parse_markdown(value), heading_offset, base_size)


def _tokens_into_docx(doc, tokens: list[Token], heading_offset: int, base_size) -> None:
    from docx.shared import Pt, Inches

    list_stack: list[bool] = []  # True = ordered
    quote_depth = 0
    heading_level: Optional[int] = None
    in_table = False
    table_header: list = []
    table_rows: list = []
    current_row: list = []
    in_thead = False
    item_first_para = False
    i = 0
    while i < len(tokens):
        tok = tokens[i]
        t = tok.type
        if t == "heading_open":
            heading_level = min(int(tok.tag[1]) + heading_offset, 9)
        elif t == "heading_close":
            heading_level = None
        elif t in ("bullet_list_open", "ordered_list_open"):
            list_stack.append(t == "ordered_list_open")
        elif t in ("bullet_list_close", "ordered_list_close"):
            if list_stack:
                list_stack.pop()
        elif t == "list_item_open":
            item_first_para = True
        elif t == "blockquote_open":
            quote_depth += 1
        elif t == "blockquote_close":
            quote_depth = max(0, quote_depth - 1)
        elif t == "table_open":
            in_table, table_header, table_rows = True, [], []
        elif t == "thead_open":
            in_thead = True
        elif t == "thead_close":
            in_thead = False
        elif t == "tr_open":
            current_row = []
        elif t == "tr_close":
            if in_thead:
                table_header = current_row
            else:
                table_rows.append(current_row)
        elif t == "table_close":
            in_table = False
            _add_docx_table(doc, [""] * len(table_header), [[""] * len(r) for r in table_rows],
                            header_tokens=table_header, row_tokens=table_rows)
        elif t == "inline":
            children = tok.children or []
            if in_table:
                current_row.append(children)
            elif heading_level is not None:
                para = doc.add_heading("", level=min(heading_level, 9))
                _write_inline(para, children)
            else:
                style = None
                if list_stack:
                    style = _list_style(doc, list_stack[-1], len(list_stack)) if item_first_para else None
                elif quote_depth:
                    style = "Quote" if _style_or_default(doc, "Quote") is not None else None
                para = doc.add_paragraph(style=style) if style else doc.add_paragraph()
                if list_stack and not item_first_para:
                    para.paragraph_format.left_indent = Inches(0.25 * (len(list_stack) + 1))
                item_first_para = False
                _write_inline(para, children, base_size=base_size)
        elif t in ("fence", "code_block"):
            para = doc.add_paragraph()
            para.paragraph_format.left_indent = Inches(0.25)
            run = para.add_run(clean_text(tok.content.rstrip("\n")))
            run.font.name = _MONO
            run.font.size = Pt(8.5)
        elif t == "hr":
            doc.add_paragraph()
        i += 1
