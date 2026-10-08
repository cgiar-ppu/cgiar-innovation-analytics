"""
create_document MCP tool — the ONLY way the IA agent produces a downloadable
file (review 2026-09-23 P0-2, sandbox 2026-09-26).

The agent used to write Word/Excel/CSV files with Bash + python-docx under
``bypassPermissions`` into a workspace shared by every user. Bash, Write and
Edit are now removed; this tool replaces that capability with a narrow,
server-side renderer:

* formats: ``docx`` (Word), ``xlsx`` (Excel), ``csv``, ``md`` (Markdown);
* input: a title, optional Markdown ``content`` and optional ``tables``
  (``[{"title", "columns", "rows"}]``);
* every file carries the mandatory AI zero-draft notice from
  ``synapsis.exporters.watermark`` (DOCX: notice box + per-page marks via
  ``apply_ai_watermark``; XLSX: a first "AI draft notice" sheet plus a printed
  header/footer on every sheet; CSV: a notice row first; Markdown: notice block
  + footer);
* the file is written to the CALLER's own area
  ``outputs/u/<owner>/<random token>/<name>.<ext>`` (``synapsis.user_files``)
  and the tool returns its absolute path, which the chat UI turns into an
  authenticated download link that only the owner can open.

The owner is resolved from the connection's verified identity
(``synapsis.auth.context``) — never from tool input.

No new dependency: DOCX uses python-docx (already required); XLSX is written as
a minimal Office Open XML package with the standard library.
"""

from __future__ import annotations

import csv
import io
import json
import re
import zipfile
from datetime import datetime, timezone
from typing import Any
from xml.sax.saxutils import escape as _xml_escape

from claude_agent_sdk import tool

from synapsis.config import logger
from synapsis.exporters.watermark import (
    PRODUCT_FOOTER,
    SOP_DISCLOSURE,
    WATERMARK_BANNER,
    apply_ai_watermark,
    export_timestamp_line,
    provenance_notice,
    snapshot_line,
    watermark_markdown,
    watermark_markdown_footer,
)
from synapsis.exporters.render import (
    markdown_into_docx,
    prepare_assistant_text,
    replace_charts_with_markdown_tables,
)
from synapsis.user_files import new_output_path
from synapsis.utils.responses import error_response, success_response

SUPPORTED_FORMATS = ("docx", "xlsx", "csv", "md")

#: Guard rails: a document is a deliverable, not a data dump.
MAX_CONTENT_CHARS = 200_000
MAX_TABLES = 20
MAX_ROWS_PER_TABLE = 20_000
MAX_COLUMNS = 100

# Excel forbids these in sheet names; 31 chars max.
_SHEET_BAD = re.compile(r"[\[\]\:\*\?\/\\]")
# XML 1.0 forbids most C0 control characters (python-docx raises on them).
_CONTROL = re.compile(r"[\x00-\x08\x0b\x0c\x0e-\x1f]")


class DocumentInputError(ValueError):
    """Raised for invalid tool input (reported back to the agent)."""


# ---------------------------------------------------------------------------
# Input normalisation
# ---------------------------------------------------------------------------

def _clean(value: Any) -> str:
    return _CONTROL.sub("", "" if value is None else str(value))


def _cell(value: Any) -> Any:
    """Keep numbers numeric (for Excel); everything else becomes clean text."""
    if isinstance(value, bool):
        return str(value)
    if isinstance(value, (int, float)):
        return value
    return _clean(value)


def normalize_tables(raw: Any) -> list[dict]:
    """Accept a list (or JSON string) of ``{title, columns, rows}`` objects."""
    if raw in (None, "", []):
        return []
    if isinstance(raw, str):
        try:
            raw = json.loads(raw)
        except json.JSONDecodeError as exc:
            raise DocumentInputError(f"'tables' is not valid JSON: {exc}") from None
    if isinstance(raw, dict):
        raw = [raw]
    if not isinstance(raw, list):
        raise DocumentInputError("'tables' must be a list of {title, columns, rows} objects")
    if len(raw) > MAX_TABLES:
        raise DocumentInputError(f"too many tables ({len(raw)} > {MAX_TABLES})")
    tables: list[dict] = []
    for i, t in enumerate(raw, 1):
        if not isinstance(t, dict):
            raise DocumentInputError(f"table {i} is not an object")
        columns = [_clean(c) for c in (t.get("columns") or [])]
        rows = t.get("rows") or []
        if not isinstance(rows, list) or any(not isinstance(r, (list, tuple)) for r in rows):
            raise DocumentInputError(f"table {i}: 'rows' must be a list of lists")
        if len(rows) > MAX_ROWS_PER_TABLE:
            raise DocumentInputError(f"table {i}: too many rows ({len(rows)} > {MAX_ROWS_PER_TABLE})")
        width = max([len(columns)] + [len(r) for r in rows]) if (columns or rows) else 0
        if width > MAX_COLUMNS:
            raise DocumentInputError(f"table {i}: too many columns ({width} > {MAX_COLUMNS})")
        if not columns:
            columns = [f"Column {n}" for n in range(1, width + 1)]
        norm_rows = [[_cell(v) for v in r] + [""] * (width - len(r)) for r in rows]
        columns = columns + [""] * (width - len(columns))
        tables.append({"title": _clean(t.get("title") or f"Table {i}"), "columns": columns, "rows": norm_rows})
    return tables


# ---------------------------------------------------------------------------
# PRMS report links for result-code columns (QA-4 D8)
# ---------------------------------------------------------------------------

#: Name of the URL column added next to a result-code column in xlsx/csv.
REPORT_URL_COLUMN = "PRMS report"

_CODE_HEADER_RE = re.compile(
    r"^\s*(?:prms\s+)?(?:result[ _\-]?code|result\s*#|code|r[ _\-]?code)s?\s*$", re.IGNORECASE
)
_URL_HEADER_RE = re.compile(r"\b(url|urls|link|links|prms report)\b", re.IGNORECASE)
# "R1003", "[R1003]", "R-1003", "[R1003](https://…)" — with the R prefix.
_R_CODE_CELL_RE = re.compile(r"^\s*\[?R-?(\d{1,7})\]?(?:\([^)\s]*\))?\s*$", re.IGNORECASE)
# Plain numbers, only trusted under a result-code header.
_NUM_CODE_CELL_RE = re.compile(r"^\s*#?(\d{1,7})\s*$")


def _cell_code(value: Any, header_is_code: bool) -> str | None:
    if isinstance(value, bool):
        return None
    if isinstance(value, int):
        return str(value) if header_is_code and value > 0 else None
    if isinstance(value, float):
        return str(int(value)) if header_is_code and value.is_integer() and value > 0 else None
    text = str(value)
    m = _R_CODE_CELL_RE.match(text)
    if m:
        return str(int(m.group(1)))
    m = _NUM_CODE_CELL_RE.match(text) if header_is_code else None
    return str(int(m.group(1))) if m else None


def result_code_column(table: dict) -> int | None:
    """Index of the table's result-code column, or None.

    A column counts when its header names a result code ("Result code",
    "Code", "R code") and its values are codes, or when every non-empty value
    is an R-prefixed code ("R1003", "[R1003]").
    """
    rows = table["rows"]
    if not rows:
        return None
    for i, col in enumerate(table["columns"]):
        values = [r[i] for r in rows if i < len(r) and str(r[i]).strip() != ""]
        if not values:
            continue
        header_is_code = bool(_CODE_HEADER_RE.match(str(col)))
        if all(_cell_code(v, header_is_code) for v in values):
            return i
    return None


def add_report_link_columns(tables: list[dict]) -> list[dict]:
    """Add a "PRMS report" URL column after each table's result-code column.

    Marc (25 Sep): a forwarded list must always carry its source URLs. The URL
    is the public per-result report from ``resolve_result_code_url`` (never a
    session-gated PRMS page); a code that is not in the snapshot gets an empty
    cell. Tables that already carry a URL/link column are left as they are.
    Returns new table dicts; each gets ``_links`` = {(row, col): url} for the
    xlsx writer (the code cell and the URL cell are both hyperlinks).
    """
    from synapsis.tools.result_code_citation import resolve_result_code_url

    out: list[dict] = []
    for t in tables:
        idx = result_code_column(t)
        if idx is None or any(_URL_HEADER_RE.search(str(c)) for c in t["columns"]):
            out.append(t)
            continue
        columns = t["columns"][: idx + 1] + [REPORT_URL_COLUMN] + t["columns"][idx + 1:]
        rows: list[list] = []
        links: dict[tuple[int, int], str] = {}
        for r_i, r in enumerate(t["rows"]):
            code = _cell_code(r[idx], True) if idx < len(r) else None
            url = (resolve_result_code_url(code) or "") if code else ""
            cell = r[idx]
            if isinstance(cell, str) and "](" in cell:  # "[R1003](url)" → "R1003"
                cell = f"R{code}" if code else cell
            rows.append(list(r[:idx]) + [cell, url] + list(r[idx + 1:]))
            if url:
                links[(r_i, idx)] = url
                links[(r_i, idx + 1)] = url
        out.append({**t, "columns": columns, "rows": rows, "_links": links})
    return out


# ---------------------------------------------------------------------------
# Markdown helpers
# ---------------------------------------------------------------------------

def _md_escape_cell(value: Any) -> str:
    return str(value).replace("|", "\\|").replace("\n", " ")


def table_to_markdown(table: dict) -> str:
    cols = table["columns"]
    lines = [f"**{table['title']}**", "", "| " + " | ".join(_md_escape_cell(c) for c in cols) + " |",
             "|" + "|".join(" --- " for _ in cols) + "|"]
    lines += ["| " + " | ".join(_md_escape_cell(v) for v in r) + " |" for r in table["rows"]]
    return "\n".join(lines)


def render_markdown(title: str, content: str, tables: list[dict], now: datetime) -> str:
    parts = [watermark_markdown(now), f"# {title}", ""]
    if content:
        # Same export pipeline as chat exports: result codes linked to their
        # public PRMS report, <chart> blocks as captioned tables (L6-02).
        parts += [replace_charts_with_markdown_tables(prepare_assistant_text(content)).strip(), ""]
    for t in tables:
        parts += [table_to_markdown(t), ""]
    parts.append(watermark_markdown_footer(now))
    return "\n".join(parts)


# ---------------------------------------------------------------------------
# CSV
# ---------------------------------------------------------------------------

def render_csv(title: str, tables: list[dict], now: datetime) -> str:
    """One table per CSV. The notice is the first row so it is always seen."""
    if len(tables) != 1:
        raise DocumentInputError("CSV needs exactly one table in 'tables' (use xlsx for several)")
    t = tables[0]
    buf = io.StringIO()
    w = csv.writer(buf)
    w.writerow([f"{WATERMARK_BANNER} — {provenance_notice(now)} {snapshot_line()} {export_timestamp_line(now)}"])
    w.writerow([f"{title} — {PRODUCT_FOOTER}"])
    w.writerow([])
    w.writerow(t["columns"])
    for r in t["rows"]:
        w.writerow(r)
    return buf.getvalue()


# ---------------------------------------------------------------------------
# DOCX (python-docx)
# ---------------------------------------------------------------------------

def _add_docx_table(doc, columns: list, rows: list) -> None:
    table = doc.add_table(rows=1, cols=max(1, len(columns)))
    table.style = "Table Grid"
    for i, c in enumerate(columns):
        cell = table.rows[0].cells[i]
        cell.text = ""
        cell.paragraphs[0].add_run(_clean(c)).bold = True
    for r in rows:
        cells = table.add_row().cells
        for i, v in enumerate(r[: len(columns)]):
            cells[i].text = _clean(v)


def render_docx(title: str, content: str, tables: list[dict], now: datetime) -> bytes:
    from docx import Document

    doc = Document()
    doc.add_heading(title, level=0)
    if content:
        # The shared chat-export renderer (review L6-02): headings, bold/italic,
        # lists, tables, clickable links (result codes linked to their public
        # PRMS report) and <chart> blocks as captioned data tables.
        markdown_into_docx(doc, prepare_assistant_text(content))
    for t in tables:
        doc.add_heading(t["title"], level=2)
        _add_docx_table(doc, t["columns"], t["rows"])
    apply_ai_watermark(doc, date=now, title=title)
    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


# ---------------------------------------------------------------------------
# XLSX (minimal Office Open XML, standard library only)
# ---------------------------------------------------------------------------

def _col_letter(n: int) -> str:
    s = ""
    while n:
        n, r = divmod(n - 1, 26)
        s = chr(65 + r) + s
    return s


def _xlsx_cell(ref: str, value: Any, style: int = 0) -> str:
    st = f' s="{style}"' if style else ""
    if isinstance(value, (int, float)) and not isinstance(value, bool):
        return f'<c r="{ref}"{st}><v>{value}</v></c>'
    text = _xml_escape(_clean(value))
    return f'<c r="{ref}" t="inlineStr"{st}><is><t xml:space="preserve">{text}</t></is></c>'


def _xlsx_sheet(rows: list[list], header_row: int | None, notice: str,
                links: dict[str, str] | None = None) -> str:
    """One worksheet. *links* maps a cell ref ("B5") to an external URL; the
    caller writes the matching relationships (rId1…) in the same order."""
    links = links or {}
    out = []
    for r_i, row in enumerate(rows, 1):
        cells = ""
        for c_i, v in enumerate(row, 1):
            ref = f"{_col_letter(c_i)}{r_i}"
            style = 1 if r_i == header_row else (2 if ref in links else 0)
            cells += _xlsx_cell(ref, v, style)
        out.append(f'<row r="{r_i}">{cells}</row>')
    hf = _xml_escape(f"&C&\"-,Bold\"{WATERMARK_BANNER}")
    ff = _xml_escape(f"&L{notice}&R&P / &N")
    hyperlinks = ""
    if links:
        hyperlinks = "<hyperlinks>" + "".join(
            f'<hyperlink ref="{ref}" r:id="rId{n}"/>' for n, ref in enumerate(links, 1)
        ) + "</hyperlinks>"
    return (
        '<?xml version="1.0" encoding="UTF-8" standalone="yes"?>'
        '<worksheet xmlns="http://schemas.openxmlformats.org/spreadsheetml/2006/main" '
        'xmlns:r="http://schemas.openxmlformats.org/officeDocument/2006/relationships">'
        f'<sheetData>{"".join(out)}</sheetData>'
        f'{hyperlinks}'
        f'<headerFooter><oddHeader>{hf}</oddHeader><oddFooter>{ff}</oddFooter></headerFooter>'
        '</worksheet>'
    )


def _xml_attr(value: str) -> str:
    return _xml_escape(value, {'"': "&quot;"})


def _xlsx_sheet_rels(links: dict[str, str]) -> str:
    return (
        '<?xml version="1.0" encoding="UTF-8" standalone="yes"?>'
        '<Relationships xmlns="http://schemas.openxmlformats.org/package/2006/relationships">'
        + "".join(
            f'<Relationship Id="rId{n}" Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink" '
            f'Target="{_xml_attr(url)}" TargetMode="External"/>'
            for n, url in enumerate(links.values(), 1)
        )
        + "</Relationships>"
    )


def _sheet_name(name: str, used: set[str]) -> str:
    base = (_SHEET_BAD.sub(" ", name).strip() or "Sheet")[:31]
    cand, n = base, 2
    while cand.lower() in used:
        suffix = f" ({n})"
        cand = base[: 31 - len(suffix)] + suffix
        n += 1
    used.add(cand.lower())
    return cand


def render_xlsx(title: str, content: str, tables: list[dict], now: datetime) -> bytes:
    if not tables:
        raise DocumentInputError("xlsx needs at least one table in 'tables'")
    footer_notice = f"{PRODUCT_FOOTER} - {WATERMARK_BANNER}"
    notice_rows = [
        [WATERMARK_BANNER],
        [provenance_notice(now)],
        [snapshot_line()],
        [export_timestamp_line(now)],
        [SOP_DISCLOSURE],
        [""],
        [title],
    ]
    if content:
        notice_rows += [[""]] + [[_clean(line)] for line in content.strip().splitlines()[:500]]
    sheets = [("AI draft notice", _xlsx_sheet(notice_rows, None, footer_notice))]
    sheet_links: list[dict[str, str]] = [{}]
    used = {"ai draft notice"}
    for t in tables:
        rows = [[t["title"]], [WATERMARK_BANNER], [], t["columns"]] + t["rows"]
        # Data rows start on sheet row 5 (title, notice, blank, header).
        links = {f"{_col_letter(c + 1)}{r + 5}": url for (r, c), url in (t.get("_links") or {}).items()}
        sheets.append((_sheet_name(t["title"], used), _xlsx_sheet(rows, 4, footer_notice, links)))
        sheet_links.append(links)

    ct_sheets = "".join(
        f'<Override PartName="/xl/worksheets/sheet{i}.xml" '
        'ContentType="application/vnd.openxmlformats-officedocument.spreadsheetml.worksheet+xml"/>'
        for i in range(1, len(sheets) + 1)
    )
    content_types = (
        '<?xml version="1.0" encoding="UTF-8" standalone="yes"?>'
        '<Types xmlns="http://schemas.openxmlformats.org/package/2006/content-types">'
        '<Default Extension="rels" ContentType="application/vnd.openxmlformats-package.relationships+xml"/>'
        '<Default Extension="xml" ContentType="application/xml"/>'
        '<Override PartName="/xl/workbook.xml" ContentType="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet.main+xml"/>'
        '<Override PartName="/xl/styles.xml" ContentType="application/vnd.openxmlformats-officedocument.spreadsheetml.styles+xml"/>'
        '<Override PartName="/docProps/core.xml" ContentType="application/vnd.openxmlformats-package.core-properties+xml"/>'
        f'{ct_sheets}</Types>'
    )
    root_rels = (
        '<?xml version="1.0" encoding="UTF-8" standalone="yes"?>'
        '<Relationships xmlns="http://schemas.openxmlformats.org/package/2006/relationships">'
        '<Relationship Id="rId1" Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/officeDocument" Target="xl/workbook.xml"/>'
        '<Relationship Id="rId2" Type="http://schemas.openxmlformats.org/package/2006/relationships/metadata/core-properties" Target="docProps/core.xml"/>'
        '</Relationships>'
    )
    wb_sheets = "".join(
        f'<sheet name="{_xml_escape(name)}" sheetId="{i}" r:id="rId{i}"/>'
        for i, (name, _) in enumerate(sheets, 1)
    )
    workbook = (
        '<?xml version="1.0" encoding="UTF-8" standalone="yes"?>'
        '<workbook xmlns="http://schemas.openxmlformats.org/spreadsheetml/2006/main" '
        'xmlns:r="http://schemas.openxmlformats.org/officeDocument/2006/relationships">'
        f'<sheets>{wb_sheets}</sheets></workbook>'
    )
    wb_rels = (
        '<?xml version="1.0" encoding="UTF-8" standalone="yes"?>'
        '<Relationships xmlns="http://schemas.openxmlformats.org/package/2006/relationships">'
        + "".join(
            f'<Relationship Id="rId{i}" Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/worksheet" Target="worksheets/sheet{i}.xml"/>'
            for i in range(1, len(sheets) + 1)
        )
        + f'<Relationship Id="rId{len(sheets) + 1}" Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/styles" Target="styles.xml"/>'
        '</Relationships>'
    )
    styles = (
        '<?xml version="1.0" encoding="UTF-8" standalone="yes"?>'
        '<styleSheet xmlns="http://schemas.openxmlformats.org/spreadsheetml/2006/main">'
        '<fonts count="3"><font><sz val="11"/><name val="Calibri"/></font>'
        '<font><b/><sz val="11"/><name val="Calibri"/></font>'
        '<font><u/><sz val="11"/><color rgb="FF0563C1"/><name val="Calibri"/></font></fonts>'
        '<fills count="2"><fill><patternFill patternType="none"/></fill><fill><patternFill patternType="gray125"/></fill></fills>'
        '<borders count="1"><border><left/><right/><top/><bottom/><diagonal/></border></borders>'
        '<cellStyleXfs count="1"><xf numFmtId="0" fontId="0" fillId="0" borderId="0"/></cellStyleXfs>'
        '<cellXfs count="3"><xf numFmtId="0" fontId="0" fillId="0" borderId="0" xfId="0"/>'
        '<xf numFmtId="0" fontId="1" fillId="0" borderId="0" xfId="0" applyFont="1"/>'
        '<xf numFmtId="0" fontId="2" fillId="0" borderId="0" xfId="0" applyFont="1"/></cellXfs>'
        '</styleSheet>'
    )
    core = (
        '<?xml version="1.0" encoding="UTF-8" standalone="yes"?>'
        '<cp:coreProperties xmlns:cp="http://schemas.openxmlformats.org/package/2006/metadata/core-properties" '
        'xmlns:dc="http://purl.org/dc/elements/1.1/" xmlns:dcterms="http://purl.org/dc/terms/" '
        'xmlns:xsi="http://www.w3.org/2001/XMLSchema-instance">'
        f'<dc:title>{_xml_escape(title)}</dc:title>'
        f'<dc:description>{_xml_escape(WATERMARK_BANNER)}</dc:description>'
        f'<dc:creator>{_xml_escape(PRODUCT_FOOTER)}</dc:creator>'
        f'<dcterms:created xsi:type="dcterms:W3CDTF">{now.strftime("%Y-%m-%dT%H:%M:%SZ")}</dcterms:created>'
        '</cp:coreProperties>'
    )
    buf = io.BytesIO()
    with zipfile.ZipFile(buf, "w", zipfile.ZIP_DEFLATED) as z:
        z.writestr("[Content_Types].xml", content_types)
        z.writestr("_rels/.rels", root_rels)
        z.writestr("docProps/core.xml", core)
        z.writestr("xl/workbook.xml", workbook)
        z.writestr("xl/_rels/workbook.xml.rels", wb_rels)
        z.writestr("xl/styles.xml", styles)
        for i, (_, xml) in enumerate(sheets, 1):
            z.writestr(f"xl/worksheets/sheet{i}.xml", xml)
            if sheet_links[i - 1]:
                z.writestr(f"xl/worksheets/_rels/sheet{i}.xml.rels", _xlsx_sheet_rels(sheet_links[i - 1]))
    return buf.getvalue()


# ---------------------------------------------------------------------------
# Orchestration (pure function, unit-testable without the SDK)
# ---------------------------------------------------------------------------

def create_document_file(
    *,
    user_id: str | None,
    title: str,
    fmt: str,
    content: str = "",
    tables: Any = None,
    filename: str | None = None,
    now: datetime | None = None,
):
    """Render and save a document in *user_id*'s area. Returns the Path."""
    fmt = (fmt or "").strip().lower().lstrip(".")
    if fmt == "markdown":
        fmt = "md"
    if fmt not in SUPPORTED_FORMATS:
        raise DocumentInputError(f"format must be one of {', '.join(SUPPORTED_FORMATS)}")
    title = _clean(title).strip()
    if not title:
        raise DocumentInputError("'title' is required")
    content = _clean(content or "")
    if len(content) > MAX_CONTENT_CHARS:
        raise DocumentInputError(f"'content' is too long ({len(content)} > {MAX_CONTENT_CHARS} chars)")
    norm = normalize_tables(tables)
    if fmt in ("xlsx", "csv"):
        # A forwarded spreadsheet keeps its sources (QA-4 D8).
        norm = add_report_link_columns(norm)
    if fmt == "docx" and not (content or norm):
        raise DocumentInputError("docx needs 'content' and/or 'tables'")
    if fmt == "md" and not (content or norm):
        raise DocumentInputError("md needs 'content' and/or 'tables'")
    now = now or datetime.now(timezone.utc)

    if fmt == "md":
        data: bytes = render_markdown(title, content, norm, now).encode("utf-8")
    elif fmt == "csv":
        data = render_csv(title, norm, now).encode("utf-8-sig")  # BOM: Excel opens UTF-8 correctly
    elif fmt == "docx":
        data = render_docx(title, content, norm, now)
    else:
        data = render_xlsx(title, content, norm, now)

    path = new_output_path(user_id, filename or title, fmt)
    path.write_bytes(data)
    return path


@tool(
    "create_document",
    "Create a downloadable file for the user: a Word document (format 'docx'), "
    "an Excel workbook ('xlsx'), a CSV ('csv') or a Markdown file ('md'). This "
    "is the ONLY way to produce a file — there is no shell or file-writing tool. "
    "Pass a 'title', optional Markdown 'content' (headings, bullet/numbered "
    "lists, **bold**/*italic*, links, pipe tables and [R<code>] citations are "
    "rendered in docx), and optional 'tables' "
    "as a JSON array of {\"title\": str, \"columns\": [str], \"rows\": [[...]]}. "
    "xlsx needs at least one table (one sheet per table); csv needs exactly one "
    "table. Give each table ONCE: either as a pipe table inside 'content' or in "
    "'tables', never both. In xlsx/csv, a table with a result-code column "
    "automatically gets a 'PRMS report' column with each result's public report "
    "URL (hyperlinked in xlsx) — do not add URLs yourself. "
    "Every number in the file MUST come from query results you actually "
    "obtained in this conversation. The file automatically carries the "
    "mandatory 'AI V0 DRAFT — REQUIRES HUMAN VALIDATION' notice and the PRMS "
    "snapshot line — do not add your own disclaimer. The tool returns the file "
    "path: include that exact path in your reply so the user gets a download "
    "link (only this user can open it).",
    {
        "type": "object",
        "properties": {
            "title": {"type": "string", "description": "Document title"},
            "format": {"type": "string", "enum": list(SUPPORTED_FORMATS)},
            "content": {"type": "string", "description": "Markdown body (optional)"},
            "tables": {
                "type": "array",
                "description": "Optional tables: [{title, columns, rows}]",
                "items": {
                    "type": "object",
                    "properties": {
                        "title": {"type": "string"},
                        "columns": {"type": "array", "items": {"type": "string"}},
                        "rows": {"type": "array", "items": {"type": "array"}},
                    },
                },
            },
            "filename": {"type": "string", "description": "Optional file name stem"},
        },
        "required": ["title", "format"],
    },
)
async def create_document(args: dict[str, Any]) -> dict[str, Any]:
    """MCP handler: render the document into the caller's own output area."""
    from synapsis.auth.context import get_current_user_id

    user_id = get_current_user_id()
    try:
        path = create_document_file(
            user_id=user_id,
            title=args.get("title") or "",
            fmt=args.get("format") or "",
            content=args.get("content") or "",
            tables=args.get("tables"),
            filename=args.get("filename"),
        )
    except DocumentInputError as exc:
        return error_response(f"create_document: {exc}")
    except Exception as exc:  # noqa: BLE001 — surface render failures to the agent
        logger.error("create_document failed: %s", exc)
        return error_response(f"create_document failed: {type(exc).__name__}")
    logger.info("create_document saved %s (%d bytes)", path.name, path.stat().st_size)
    return success_response(
        "Document created.\n\n"
        f"**File:** `{path}`\n\n"
        "Paste this exact path into your reply, exactly as shown and with NOTHING "
        "in front of it (no `sandbox:`, `file://` or other scheme), either as plain "
        "text or as the target of a Markdown link whose text is the file name, e.g. "
        f"[{path.name}]({path}). The chat turns it into a download button that shows "
        "only the file name. It carries the AI zero-draft notice."
    )
