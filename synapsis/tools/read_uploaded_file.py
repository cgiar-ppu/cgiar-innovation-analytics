"""
read_uploaded_file MCP tool — read a file the CALLING user uploaded (2026-09-26).

Lane A's sandbox removed Bash/Python from the agent, and with them the only
way it had to open binary uploads (``.xlsx``, ``.docx``, ``.pdf``): ``Read``
works for plain text only. This narrow tool restores that capability without
giving back a shell:

* **input**: the upload's path exactly as the chat message shows it
  (``/workspace/uploads/u/<owner>/<name>``) or just its file name;
* **owner check**: the resolved real path must lie inside the caller's OWN
  uploads area (``synapsis.user_files.owner_area("uploads", <caller>)``), the
  same rule as the Read-path hook in ``hooks/sandbox.py``; the caller comes
  from the connection's verified identity (``synapsis.auth.context``), never
  from tool input. Other users' uploads, outputs, references, hidden files and
  anything else are refused;
* **output**: text / Markdown — ``.xlsx``/``.xlsm``/``.csv``/``.tsv`` become one
  Markdown table per sheet with its row and column counts (at most
  ``MAX_ROWS`` rows and ``MAX_COLUMNS`` columns per sheet, with a truncation
  note); ``.docx`` (paragraphs + tables), ``.pdf`` (text per page) and plain
  text files are returned as text capped at ``MAX_CHARS`` characters.

Parsers: openpyxl (+ defusedxml, pinned in requirements.txt), python-docx,
pypdf and the standard library. Files over ``MAX_FILE_BYTES`` and archives
that inflate beyond ``MAX_UNZIPPED_BYTES`` are refused.
"""

from __future__ import annotations

import asyncio
import csv
import io
import os
import zipfile
from datetime import date, datetime, time
from pathlib import Path, PurePosixPath
from typing import Any, Iterable

from claude_agent_sdk import tool

from synapsis.config import logger
from synapsis.utils.responses import error_response, success_response

MAX_ROWS = 200
MAX_COLUMNS = 50
MAX_CHARS = 50_000
MAX_SHEETS = 20
MAX_PDF_PAGES = 300
MAX_FILE_BYTES = 25 * 1024 * 1024
MAX_UNZIPPED_BYTES = 200 * 1024 * 1024
#: Rows counted per sheet beyond the displayed ones (bounds the work on huge sheets).
MAX_COUNTED_ROWS = 1_000_000

TABLE_TYPES = {".xlsx", ".xlsm", ".csv", ".tsv"}
TEXT_TYPES = {".txt", ".md", ".markdown", ".json", ".xml", ".html", ".htm", ".log", ".yaml", ".yml"}
SUPPORTED = TABLE_TYPES | TEXT_TYPES | {".docx", ".pdf"}


class UploadAccessError(PermissionError):
    """The path is not one of the caller's own uploads."""


class UploadReadError(ValueError):
    """The file is the caller's, but cannot be read (type, size, corrupt)."""


# ---------------------------------------------------------------------------
# Ownership
# ---------------------------------------------------------------------------

def resolve_own_upload(raw: str | None, user_id: str | None) -> Path:
    """Resolve *raw* to a file inside *user_id*'s uploads area, or raise.

    A bare file name is looked up in the caller's uploads folder. Symlinks
    are resolved first, so a link pointing elsewhere is refused.
    """
    from synapsis.user_files import owner_area

    if not raw or not str(raw).strip():
        raise UploadAccessError("a file path or name is required")
    if not user_id:
        raise UploadAccessError("no signed-in user")
    raw = str(raw).strip()
    if raw.startswith("file://"):
        raw = raw[len("file://"):]
    root = Path(os.path.realpath(owner_area("uploads", user_id)))
    candidate = Path(os.path.expanduser(raw))
    if not candidate.is_absolute():
        if len(PurePosixPath(raw).parts) != 1:
            raise UploadAccessError("give the full path shown in the chat, or just the file name")
        candidate = root / raw
    real = Path(os.path.realpath(candidate))
    if real != root and root not in real.parents:
        raise UploadAccessError("that file is not one of this user's own uploads")
    rel_parts = real.relative_to(root).parts
    if not rel_parts or any(p.startswith(".") for p in rel_parts):
        raise UploadAccessError("that file is not one of this user's own uploads")
    if not real.is_file():
        raise UploadReadError("no such uploaded file (check the exact name in the chat)")
    return real


# ---------------------------------------------------------------------------
# Rendering helpers
# ---------------------------------------------------------------------------

def _cell_text(value: Any) -> str:
    if value is None:
        return ""
    if isinstance(value, bool):
        return "TRUE" if value else "FALSE"
    if isinstance(value, float):
        if value.is_integer():
            return str(int(value))
        return repr(value)
    if isinstance(value, datetime):
        return value.isoformat(sep=" ", timespec="seconds").replace(" 00:00:00", "")
    if isinstance(value, (date, time)):
        return value.isoformat()
    text = str(value)
    return text.replace("\r\n", " ").replace("\n", " ").replace("|", "\\|").strip()


def _markdown_table(rows: list[list[str]], width: int) -> str:
    if not rows:
        return "_(empty)_"
    header = rows[0] + [""] * (width - len(rows[0]))
    header = [h or f"Column {i + 1}" for i, h in enumerate(header)]
    lines = ["| " + " | ".join(header) + " |", "|" + "|".join(" --- " for _ in header) + "|"]
    for r in rows[1:]:
        r = r + [""] * (width - len(r))
        lines.append("| " + " | ".join(r) + " |")
    return "\n".join(lines)


def _table_block(name: str, rows_iter: Iterable[Iterable[Any]]) -> str:
    """One sheet → heading, counts, capped Markdown table, truncation notes."""
    shown: list[list[str]] = []
    total_rows = 0
    width = 0
    for raw in rows_iter:
        cells = [_cell_text(v) for v in raw]
        while cells and cells[-1] == "":
            cells.pop()
        if not cells:
            continue  # skip fully empty rows
        total_rows += 1
        width = max(width, len(cells))
        if len(shown) <= MAX_ROWS:  # header + MAX_ROWS data rows
            shown.append(cells)
        if total_rows >= MAX_COUNTED_ROWS:
            break
    data_rows = max(0, total_rows - 1)
    shown_width = min(width, MAX_COLUMNS)
    shown = [r[:shown_width] for r in shown]
    parts = [f"### {name}", "",
             f"{data_rows:,} data row(s) + 1 header row, {width} column(s)"
             + (" (counting stopped at the limit)" if total_rows >= MAX_COUNTED_ROWS else "") + ".", ""]
    parts.append(_markdown_table(shown, shown_width) if total_rows else "_(sheet is empty)_")
    notes = []
    if data_rows > MAX_ROWS:
        notes.append(f"only the first {MAX_ROWS} of {data_rows:,} data rows are shown")
    if width > MAX_COLUMNS:
        notes.append(f"only the first {MAX_COLUMNS} of {width} columns are shown")
    if notes:
        parts += ["", f"_Truncated: {'; '.join(notes)}. Ask the user for a filtered or smaller extract if the rest matters._"]
    return "\n".join(parts)


def _check_zip(path: Path) -> None:
    try:
        with zipfile.ZipFile(path) as z:
            total = sum(i.file_size for i in z.infolist())
    except zipfile.BadZipFile:
        raise UploadReadError("the file is not a valid Office document (corrupt or wrong extension)") from None
    if total > MAX_UNZIPPED_BYTES:
        raise UploadReadError("the document is too large to read once unpacked")


def _cap(text: str) -> str:
    if len(text) <= MAX_CHARS:
        return text
    return text[:MAX_CHARS] + f"\n\n_Truncated: showing the first {MAX_CHARS:,} of {len(text):,} characters._"


# ---------------------------------------------------------------------------
# Readers
# ---------------------------------------------------------------------------

def _read_xlsx(path: Path) -> str:
    _check_zip(path)
    try:
        import openpyxl
    except ImportError:  # pragma: no cover - pinned in requirements.txt
        raise UploadReadError("Excel reading is not available on this server") from None
    try:
        wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    except Exception as exc:  # noqa: BLE001 — corrupt / encrypted workbooks
        raise UploadReadError(f"the workbook could not be opened ({type(exc).__name__})") from None
    try:
        names = wb.sheetnames
        blocks = [f"Workbook with {len(names)} sheet(s): {', '.join(names)}."]
        for name in names[:MAX_SHEETS]:
            ws = wb[name]
            blocks.append(_table_block(f"Sheet: {name}", ws.iter_rows(values_only=True)))
        if len(names) > MAX_SHEETS:
            blocks.append(f"_Truncated: only the first {MAX_SHEETS} of {len(names)} sheets are shown._")
    finally:
        wb.close()
    return "\n\n".join(blocks)


def _read_delimited(path: Path) -> str:
    raw = path.read_bytes()
    text = raw.decode("utf-8-sig", errors="replace")
    if path.suffix.lower() == ".tsv":
        delimiter = "\t"
    else:
        try:
            delimiter = csv.Sniffer().sniff(text[:20_000], delimiters=",;\t|").delimiter
        except csv.Error:
            delimiter = ","
    reader = csv.reader(io.StringIO(text), delimiter=delimiter)
    return _table_block(f"Table: {path.name}", reader)


def _read_docx(path: Path) -> str:
    _check_zip(path)
    try:
        from docx import Document
    except ImportError:  # pragma: no cover - pinned in requirements.txt
        raise UploadReadError("Word reading is not available on this server") from None
    try:
        doc = Document(str(path))
    except Exception as exc:  # noqa: BLE001
        raise UploadReadError(f"the Word document could not be opened ({type(exc).__name__})") from None
    from docx.table import Table
    from docx.text.paragraph import Paragraph

    parts: list[str] = []
    size = 0
    n_table = 0
    for child in doc.element.body.iterchildren():
        tag = child.tag.rsplit("}", 1)[-1]
        if tag == "p":
            para = Paragraph(child, doc)
            text = para.text.strip()
            if not text:
                continue
            style = (para.style.name if para.style is not None else "") or ""
            if style.startswith("Heading"):
                level = style.replace("Heading", "").strip()
                text = "#" * (int(level) if level.isdigit() else 2) + " " + text
            elif style.startswith("List"):
                text = "- " + text
            parts.append(text)
        elif tag == "tbl":
            n_table += 1
            table = Table(child, doc)
            parts.append(_table_block(f"Table {n_table}", ([c.text for c in r.cells] for r in table.rows)))
        size += len(parts[-1]) if parts else 0
        if size > MAX_CHARS * 2:
            break
    return _cap("\n\n".join(parts) or "_(the document has no text)_")


def _read_pdf(path: Path) -> str:
    try:
        from pypdf import PdfReader
    except ImportError:  # pragma: no cover - pinned in requirements.txt
        raise UploadReadError("PDF reading is not available on this server") from None
    try:
        reader = PdfReader(str(path))
        if reader.is_encrypted:
            raise UploadReadError("the PDF is password-protected")
        pages = reader.pages
        n = len(pages)
    except UploadReadError:
        raise
    except Exception as exc:  # noqa: BLE001
        raise UploadReadError(f"the PDF could not be opened ({type(exc).__name__})") from None
    parts = [f"PDF with {n} page(s)."]
    size = 0
    for i in range(min(n, MAX_PDF_PAGES)):
        try:
            text = (pages[i].extract_text() or "").strip()
        except Exception:  # noqa: BLE001 — one bad page must not lose the rest
            text = "_(text of this page could not be extracted)_"
        parts.append(f"--- Page {i + 1} ---\n{text or '_(no extractable text; the page may be an image)_'}")
        size += len(parts[-1])
        if size > MAX_CHARS:
            break
    if n > MAX_PDF_PAGES:
        parts.append(f"_Only the first {MAX_PDF_PAGES} pages were read._")
    return _cap("\n\n".join(parts))


def _read_text(path: Path) -> str:
    with path.open("rb") as fh:
        raw = fh.read(MAX_CHARS * 4 + 1)
    return _cap(raw.decode("utf-8-sig", errors="replace"))


def read_upload_text(path: Path) -> str:
    """Render an (already owner-checked) upload as text/Markdown. Blocking."""
    ext = path.suffix.lower()
    size = path.stat().st_size
    if ext == ".xls":
        raise UploadReadError("old .xls workbooks are not supported; ask the user to save it as .xlsx and upload again")
    if ext not in SUPPORTED:
        raise UploadReadError(
            f"'{ext or 'no extension'}' files are not supported (supported: "
            + ", ".join(sorted(SUPPORTED)) + ")"
        )
    if size > MAX_FILE_BYTES:
        raise UploadReadError(f"the file is too large to read ({size / 1_048_576:.1f} MB > {MAX_FILE_BYTES // 1_048_576} MB)")
    if ext in (".xlsx", ".xlsm"):
        body = _read_xlsx(path)
    elif ext in (".csv", ".tsv"):
        body = _read_delimited(path)
    elif ext == ".docx":
        body = _read_docx(path)
    elif ext == ".pdf":
        body = _read_pdf(path)
    else:
        body = _read_text(path)
    return f"**Uploaded file:** {path.name} ({size:,} bytes)\n\n{body}"


def read_uploaded_file_for(user_id: str | None, raw_path: str | None) -> str:
    """Owner-check *raw_path* for *user_id* and return its text. Raises on refusal."""
    return read_upload_text(resolve_own_upload(raw_path, user_id))


# ---------------------------------------------------------------------------
# MCP tool
# ---------------------------------------------------------------------------

@tool(
    "read_uploaded_file",
    "Read a file THIS user uploaded to the chat: Excel (.xlsx/.xlsm), CSV/TSV, "
    "Word (.docx), PDF or plain text. Pass 'path' exactly as shown in the "
    "upload message (or just the file name). Tables come back as Markdown "
    f"tables per sheet with row/column counts (first {MAX_ROWS} rows and "
    f"{MAX_COLUMNS} columns per sheet); documents as text (first "
    f"{MAX_CHARS:,} characters). Only the calling user's own uploads can be read.",
    {
        "type": "object",
        "properties": {
            "path": {"type": "string", "description": "Upload path from the chat message, or the file name"},
        },
        "required": ["path"],
    },
)
async def read_uploaded_file(args: dict[str, Any]) -> dict[str, Any]:
    """MCP handler: owner-checked, bounded text rendering of one upload."""
    from synapsis.auth.context import get_current_user_id

    user_id = get_current_user_id()
    raw = args.get("path") or args.get("file_path") or args.get("name")
    try:
        text = await asyncio.to_thread(read_uploaded_file_for, user_id, raw)
    except UploadAccessError as exc:
        logger.info("read_uploaded_file denied: %s", exc)
        return error_response(f"read_uploaded_file: {exc}.")
    except UploadReadError as exc:
        return error_response(f"read_uploaded_file: {exc}.")
    except Exception as exc:  # noqa: BLE001 — never leak a traceback to the model
        logger.error("read_uploaded_file failed: %s", type(exc).__name__)
        return error_response(f"read_uploaded_file failed: {type(exc).__name__}")
    return success_response(text)
