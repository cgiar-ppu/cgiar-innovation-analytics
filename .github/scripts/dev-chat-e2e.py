"""DEV-only end-to-end chat check with real questions (opt-in, request-driven).

Runs INSIDE the DEV container (base64-exec'd via SSM ``docker exec`` by
dev-chat-e2e.yml), the same pattern as dev-opus55-check.py / dev-qa-smoke.py:
synthetic identities minted in the container (never printed), no real user
account, no auth bypass.

Input  (env IA_E2E_REQUEST, base64 JSON):
  {"questions": [{"id": "q1", "text": "...",
                  "role": "researcher" | "external" (optional, default researcher),
                  "expect": "declined_financial" | "answered" (optional)}],
   "public_key_pem": "...",
   "model": "claude-sonnet-5" (optional), "max_cost_usd": 5.0, "blob": "/tmp/x.b64"}

``role`` picks the synthetic identity that asks the question (one per role).
``expect`` adds a pass/fail check to the summary line (booleans only):
  * declined_financial — the external-view refusal: no USD/currency amount in
    the answer, the answer says the detail is not available in this external
    view, and no tool call returned budget figures;
  * answered — a normal answer: no error, at least one successful PRMS tool
    call, and no external-view refusal wording;
  * answered_with_amount — ``answered`` plus a currency amount in the answer
    (full-access control: budgets are still readable);
  * no_amounts — whatever the model answers (quote, refusal, file), no
    currency amount appears in the answer, in any tool input or result, or in
    any file the chat produced (xlsx/docx/csv/md cells are read back).
The workflow fails the run (after saving the encrypted transcript) when a
check fails.
Output:
  * stdout: ONE line {"dev_chat_e2e": {...}} with non-sensitive stats only
    (costs, turns, tool names, error codes). The repository is public.
  * the full transcripts (answers, tool inputs/outputs, detail=full exports,
    dashboard KPIs) are compressed and encrypted to the requester's RSA public
    key (RSA-OAEP-SHA256 wrapping an AES-256-GCM key) and written to ``blob``;
    the workflow reads it out in chunks, publishes it as a short-lived artifact
    and deletes the file. Only the holder of the private key can read it.
Every chat it creates is titled "[QA] ..." and deleted at the end.
"""
import asyncio, base64, json, os, re, time, zlib

import httpx, websockets
from cryptography.hazmat.primitives import hashes, serialization
from cryptography.hazmat.primitives.asymmetric import padding
from cryptography.hazmat.primitives.ciphers.aead import AESGCM
from synapsis.auth.tokens import create_access_token

BASE = os.environ.get('IA_E2E_BASE', 'http://localhost:7780')
REQ = json.loads(base64.b64decode(os.environ['IA_E2E_REQUEST']))
TOOL_RESULT_CAP = 6000
EXPORT_CAP = 120_000
ROLES = ('researcher', 'external')
# A currency amount in an answer, e.g. "USD 1.2 million", "$450,000", "US$ 3m",
# "4.5 million euros". Deliberately independent of (and looser than) the app's
# own redaction pattern.
_MONEY = re.compile(r'(?:USD|US\$|\$|EUR|Euros?|€|GBP|£)\s?\d'
                    r'|\d[\d,.]*\s?(?:million|billion|bn|m|k)?\s?(?:USD|US dollars|dollars|EUR|euros|GBP)\b', re.I)
_FILE_RE = re.compile(r"\*\*File:\*\* `([^`]+)`")
# The external-view refusal: "not available in this external view", or the model's
# paraphrases seen on DEV ("this external partner view ... aren't available").
_EXTERNAL_VIEW_PHRASE = re.compile(r"external(?: partner)? view", re.I)
_EXTERNAL_VIEW = re.compile(r"external(?: partner)? view|(?:not|n't|n’t) available", re.I)
_PRMS_TOOLS = ('prms_query', 'prms_search')


def check(rec, expect):
    """Pass/fail for a question's ``expect`` (non-sensitive booleans only)."""
    answer = rec.get('answer') or ''
    tools = rec.get('tools', [])
    prms = [t for t in tools if any(n in str(t.get('tool')) for n in _PRMS_TOOLS)]
    if expect == 'declined_financial':
        tool_money = any(_MONEY.search(str(t.get('result') or '')) for t in tools)
        parts = {'no_error': not rec.get('error'), 'no_amount_in_answer': not _MONEY.search(answer),
                 'says_not_available': bool(_EXTERNAL_VIEW.search(answer)),
                 'no_amount_from_tools': not tool_money}
    elif expect in ('answered', 'answered_with_amount'):
        parts = {'no_error': not rec.get('error'), 'has_answer': len(answer.strip()) > 40,
                 'prms_tool_ok': any(not t.get('is_error') for t in prms),
                 'no_refusal_wording': not _EXTERNAL_VIEW_PHRASE.search(answer)}
        if expect == 'answered_with_amount':
            parts['amount_in_answer'] = bool(_MONEY.search(answer))
    elif expect == 'no_amounts':
        files = rec.get('files', [])
        parts = {'no_error': not rec.get('error'), 'has_answer': len(answer.strip()) > 20,
                 'no_amount_in_answer': not _MONEY.search(answer),
                 'no_amount_in_tool_inputs': not any(_MONEY.search(json.dumps(t.get('input'))) for t in tools),
                 'no_amount_from_tools': not any(t.get('result_amounts') or _MONEY.search(str(t.get('result') or ''))
                                                 for t in tools),
                 'no_amount_in_files': not any(f.get('amounts') for f in files),
                 'files_readable': all(f.get('read') for f in files)}
        rec['files_seen'] = len(files)
    else:
        return None
    return {'expect': expect, 'pass': all(parts.values()), **parts}


async def ask(token, question, model, h):
    rec = {'id': question['id'], 'question': question['text'], 'events': [], 'tools': []}
    async with websockets.connect(BASE.replace('http', 'ws') + '/ws/chat?token=' + token,
                                  open_timeout=30, max_size=16 * 1024 * 1024) as ws:
        async def until(kind, timeout):
            got = []
            async with asyncio.timeout(timeout):
                while True:
                    e = json.loads(await ws.recv()); got.append(e)
                    if e.get('type') == 'error':
                        rec['error'] = {'code': e.get('code'), 'message': str(e.get('message', ''))[:500]}
                        return got
                    if e.get('type') == kind:
                        return got
        await ws.send(json.dumps({'type': 'new_session'}))
        sid = (await until('session', 90))[-1]['session_id']
        rec['session_id'] = sid
        async with httpx.AsyncClient() as c:
            await c.patch(BASE + '/api/sessions/' + sid, headers=h, json={'title': f"[QA] E2E {question['id']}"})
        if model:
            await ws.send(json.dumps({'type': 'switch_model', 'model': model}))
            await until('model_switched', 90)
        t0 = time.time()
        await ws.send(json.dumps({'message': question['text']}))
        try:
            events = await until('result', 1200)
        except TimeoutError:
            events = []; rec['error'] = {'code': 'client_timeout'}
        rec['seconds'] = round(time.time() - t0)
    rec['answer'] = ''.join(e.get('content', '') for e in events if e.get('type') == 'text')
    for e in events:
        if e.get('type') == 'tool_use':
            rec['tools'].append({'tool': e.get('tool'), 'input': e.get('input'), 'id': e.get('tool_use_id')})
        elif e.get('type') == 'tool_result':
            for t in rec['tools']:
                if t['id'] == e.get('tool_use_id'):
                    content = e.get('content')
                    content = content if isinstance(content, str) else json.dumps(content)
                    t['result'] = content[:TOOL_RESULT_CAP]; t['result_len'] = len(content)
                    t['result_amounts'] = len(_MONEY.findall(content))  # over the FULL result, not the cap
                    t['is_error'] = e.get('is_error')
        elif e.get('type') == 'result':
            rec['result'] = {k: e.get(k) for k in ('estimated_cost', 'turns', 'model', 'duration_ms') if k in e}
    rec['files'] = [read_file(p) for t in rec['tools'] for p in _FILE_RE.findall(str(t.get('result') or ''))]
    async with httpx.AsyncClient(timeout=120) as c:
        r = await c.get(BASE + '/api/export/' + rec['session_id'],
                        params={'format': 'md', 'detail': 'full', 'token': token})
        rec['export_full_md'] = r.text[:EXPORT_CAP] if r.status_code == 200 else f'export status {r.status_code}'
    return rec


def read_file(path):
    """Text of a file the chat produced (all cells / paragraphs), its currency
    amounts count, then the file is deleted (synthetic QA owner)."""
    info = {'name': os.path.basename(path), 'read': False}
    try:
        if path.endswith('.xlsx'):
            import openpyxl
            wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
            text = '\n'.join(' | '.join('' if v is None else str(v) for v in row)
                             for ws in wb.worksheets for row in ws.iter_rows(values_only=True))
        elif path.endswith('.docx'):
            import docx
            d = docx.Document(path)
            text = '\n'.join([p.text for p in d.paragraphs]
                             + [c.text for t in d.tables for r in t.rows for c in r.cells])
        else:
            text = open(path, encoding='utf-8', errors='replace').read()
        info.update(read=True, chars=len(text), amounts=len(_MONEY.findall(text)),
                    placeholders=text.count('[amount withheld in external view]'), text=text[:EXPORT_CAP])
    except Exception as e:
        info['error'] = repr(e)[:200]
    try:
        os.remove(path); info['deleted'] = True
    except OSError:
        info['deleted'] = False
    return info


async def cleanup(out, headers):
    # Clean-up: delete the [QA] chats. A DELETE can race the server's orphan reaper
    # (client_registry.cleanup_session_client KeyError -> 500, seen 14 Oct), so
    # tolerate errors and retry once; the result is reported, never fatal.
    deleted = []
    await asyncio.sleep(3)
    async with httpx.AsyncClient(timeout=60) as c:
        for rec in out['questions']:
            if not rec.get('session_id'):
                continue
            h = headers[rec.get('role', 'researcher')]
            status = None
            for attempt in range(2):
                try:
                    status = (await c.delete(BASE + '/api/sessions/' + rec['session_id'], headers=h)).status_code
                except httpx.HTTPError as e:
                    status = type(e).__name__
                if status in (200, 204, 404):
                    break
                await asyncio.sleep(3)
            deleted.append(status)
        # Optional: remove [QA] chats left by an earlier run whose clean-up failed.
        # The earlier synthetic owner id is 'dev-e2e-qa-<epoch>'; try each epoch in
        # the given range — DELETE is owner-only, so only the real owner succeeds.
        prev = REQ.get('cleanup_previous') or {}
        leftovers = {}
        if prev.get('session_ids'):
            lo, hi = prev['owner_epoch_range']
            assert 0 <= hi - lo <= 120, 'owner_epoch_range too wide'
            for sid in prev['session_ids']:
                leftovers[sid] = 'not found'
                for ts in range(lo, hi + 1):
                    tok = create_access_token('dev-e2e-qa-' + str(ts), 'Synthetic E2E researcher QA', 'researcher',
                                              auth_source='sso', lifetime_seconds=120)
                    try:
                        st = (await c.delete(BASE + '/api/sessions/' + sid,
                                             headers={'Authorization': 'Bearer ' + tok})).status_code
                    except httpx.HTTPError as e:
                        st = type(e).__name__
                    if st == 200:
                        leftovers[sid] = 'deleted'; break
    return {'this_run': deleted, 'previous': leftovers}


async def main():
    stamp = str(int(time.time()))
    owner = 'dev-e2e-qa-' + stamp
    token = create_access_token(owner, 'Synthetic E2E researcher QA', 'researcher',
                                auth_source='sso', lifetime_seconds=3600)
    h = {'Authorization': 'Bearer ' + token}
    # One synthetic identity per role; the external one exercises the partner
    # view (no budget / financial data). Never an admin.
    tokens = {'researcher': token,
              'external': create_access_token('dev-e2e-qa-ext-' + stamp, 'Synthetic E2E external QA', 'external',
                                              auth_source='sso', lifetime_seconds=3600)}
    headers = {r: {'Authorization': 'Bearer ' + t} for r, t in tokens.items()}
    out = {'started_utc': time.strftime('%Y-%m-%dT%H:%M:%SZ', time.gmtime()), 'questions': []}
    async with httpx.AsyncClient(timeout=120) as c:
        out['health'] = (await c.get(BASE + '/api/health')).json()
        out['config_data'] = {k: v for k, v in (await c.get(BASE + '/api/config', headers=h)).json().items()
                              if 'snapshot' in k or 'prms' in k or k in ('default_model', 'model')}
        out['kpis'] = {}
        for label, params in (('all_years', {}), ('2025', {'years': '2025'})):
            r = await c.get(BASE + '/api/dashboard/prms-stats', headers=h, params=params)
            body = r.json() if r.status_code == 200 else {'status': r.status_code}
            out['kpis'][label] = {'kpis': body.get('kpis'), 'kpi_errors': body.get('kpi_errors'),
                                  'snapshot': body.get('snapshot')}
        try:
            r = await c.get(BASE + '/api/prms/snapshot', headers=h)
            out['prms_snapshot'] = r.json() if r.status_code == 200 else {'status': r.status_code}
        except Exception as e:
            out['prms_snapshot'] = {'error': repr(e)[:200]}
    spent = 0.0
    for q in REQ['questions']:
        if spent >= float(REQ.get('max_cost_usd', 5)):
            out['questions'].append({'id': q['id'], 'skipped': 'budget'}); continue
        role = q.get('role', 'researcher')
        assert role in ROLES, 'unknown role ' + repr(role)
        rec = await ask(tokens[role], q, REQ.get('model'), headers[role])
        rec['role'] = role
        rec['check'] = check(rec, q.get('expect'))
        spent += float((rec.get('result') or {}).get('estimated_cost') or 0)
        out['questions'].append(rec)
    pub = serialization.load_pem_public_key(REQ['public_key_pem'].encode())
    key = AESGCM.generate_key(256); nonce = os.urandom(12)
    ct = AESGCM(key).encrypt(nonce, zlib.compress(json.dumps(out).encode(), 9), None)
    wrapped = pub.encrypt(key, padding.OAEP(mgf=padding.MGF1(hashes.SHA256()), algorithm=hashes.SHA256(), label=None))
    blob = json.dumps({'k': base64.b64encode(wrapped).decode(), 'n': base64.b64encode(nonce).decode(),
                       'c': base64.b64encode(ct).decode()})
    with open(REQ['blob'], 'w') as f:  # written BEFORE clean-up so a clean-up error cannot lose it
        f.write(blob)
    try:
        out_cleanup = await cleanup(out, headers)
    except Exception as e:  # never lose the transcript summary over clean-up
        out_cleanup = {'error': repr(e)[:300]}
    stats = {'model': REQ.get('model') or 'default', 'cost_usd': round(spent, 4), 'blob_chars': len(blob),
             'kpi_errors': {k: bool(v.get('kpi_errors')) for k, v in out['kpis'].items()},
             'questions': [{'id': r['id'], 'seconds': r.get('seconds'), 'turns': (r.get('result') or {}).get('turns'),
                            'cost': (r.get('result') or {}).get('estimated_cost'),
                            'error': (r.get('error') or {}).get('code'), 'skipped': r.get('skipped'),
                            'tools': [t['tool'] for t in r.get('tools', [])],
                            'tool_errors': sum(1 for t in r.get('tools', []) if t.get('is_error')),
                            'role': r.get('role'), 'check': r.get('check'),
                            'files': [{k: f.get(k) for k in ('read', 'amounts', 'placeholders', 'deleted')}
                                      for f in r.get('files', [])]}
                           for r in out['questions']],
             'cleanup': out_cleanup}
    stats['checks_failed'] = [r['id'] for r in out['questions'] if (r.get('check') or {}).get('pass') is False]
    print(json.dumps({'dev_chat_e2e': stats}))


asyncio.run(main())
