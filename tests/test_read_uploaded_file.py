"""read_uploaded_file: the owner-checked reader for users' uploads (Lane C addendum).

The sandbox removed Python from the agent, so it could no longer open .xlsx /
.docx / .pdf uploads. This narrow MCP tool reads ONLY the calling user's own
uploads and returns bounded text / Markdown.
"""

from __future__ import annotations

import asyncio
import io
import os
from unittest.mock import patch

import pytest

ALICE = "alice@cgiar.org"
BOB = "bob@cgiar.org"


def _module():
    # ``synapsis.tools.read_uploaded_file`` the ATTRIBUTE is the tool object
    # (tools/__init__.py imports it); the module lives in sys.modules.
    import importlib

    return importlib.import_module("synapsis.tools.read_uploaded_file")


@pytest.fixture
def ws(tmp_path):
    ws = tmp_path / "workspace"
    ws.mkdir()
    with patch("synapsis.config.WORKSPACE", ws):
        yield ws


def _upload(ws, user, name, data: bytes | str):
    from synapsis.user_files import owner_area

    d = owner_area("uploads", user, ws)
    d.mkdir(parents=True, exist_ok=True)
    p = d / name
    p.write_bytes(data.encode() if isinstance(data, str) else data)
    return p


def _xlsx_bytes(sheets: dict[str, list[list]]) -> bytes:
    import openpyxl

    wb = openpyxl.Workbook()
    wb.remove(wb.active)
    for name, rows in sheets.items():
        ws = wb.create_sheet(name)
        for r in rows:
            ws.append(r)
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def _docx_bytes() -> bytes:
    from docx import Document

    doc = Document()
    doc.add_heading("Kenya scaling plan", level=1)
    doc.add_paragraph("We aim to scale three innovations.")
    t = doc.add_table(rows=2, cols=2)
    t.rows[0].cells[0].text, t.rows[0].cells[1].text = "Innovation", "IRL"
    t.rows[1].cells[0].text, t.rows[1].cells[1].text = "Drought-tolerant maize", "8"
    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


def _pdf_bytes(text_pages: list[str]) -> bytes:
    """A tiny valid PDF with one text line per page (no extra dependency)."""
    objs = []
    kids = []
    n_pages = len(text_pages)
    # 1: catalog, 2: pages, 3: font, then page/content pairs
    for i, text in enumerate(text_pages):
        page_id = 4 + 2 * i
        kids.append(f"{page_id} 0 R")
    objs.append("<< /Type /Catalog /Pages 2 0 R >>")
    objs.append(f"<< /Type /Pages /Kids [{' '.join(kids)}] /Count {n_pages} >>")
    objs.append("<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>")
    for i, text in enumerate(text_pages):
        content_id = 5 + 2 * i
        stream = f"BT /F1 12 Tf 72 720 Td ({text}) Tj ET"
        objs.append(f"<< /Type /Page /Parent 2 0 R /MediaBox [0 0 612 792] "
                    f"/Resources << /Font << /F1 3 0 R >> >> /Contents {content_id} 0 R >>")
        objs.append(f"<< /Length {len(stream)} >>\nstream\n{stream}\nendstream")
    out = io.BytesIO()
    out.write(b"%PDF-1.4\n")
    offsets = []
    for i, body in enumerate(objs, 1):
        offsets.append(out.tell())
        out.write(f"{i} 0 obj\n{body}\nendobj\n".encode())
    xref = out.tell()
    out.write(f"xref\n0 {len(objs) + 1}\n0000000000 65535 f \n".encode())
    for off in offsets:
        out.write(f"{off:010d} 00000 n \n".encode())
    out.write(f"trailer\n<< /Size {len(objs) + 1} /Root 1 0 R >>\nstartxref\n{xref}\n%%EOF\n".encode())
    return out.getvalue()


# ---------------------------------------------------------------------------
# Owner can read their own uploads
# ---------------------------------------------------------------------------

def test_owner_reads_own_xlsx_as_markdown_tables_per_sheet(ws):
    from synapsis.tools.read_uploaded_file import read_uploaded_file_for

    p = _upload(ws, ALICE, "portfolio.xlsx", _xlsx_bytes({
        "Kenya": [["Innovation", "IRL", "Share"], ["Maize", 8, 0.56], ["Beans", None, "n/a"]],
        "Notes": [["Source: PRMS"]],
    }))
    text = read_uploaded_file_for(ALICE, str(p))
    assert "Workbook with 2 sheet(s): Kenya, Notes." in text
    assert "### Sheet: Kenya" in text and "2 data row(s) + 1 header row, 3 column(s)." in text
    assert "| Innovation | IRL | Share |" in text and "| Maize | 8 | 0.56 |" in text
    assert "### Sheet: Notes" in text


def test_owner_reads_csv_docx_pdf_and_text(ws):
    from synapsis.tools.read_uploaded_file import read_uploaded_file_for

    csv_p = _upload(ws, ALICE, "data.csv", "Year;Count\n2024;1016\n2025;1185\n")
    out = read_uploaded_file_for(ALICE, str(csv_p))
    assert "| Year | Count |" in out and "| 2025 | 1185 |" in out

    docx_p = _upload(ws, ALICE, "plan.docx", _docx_bytes())
    out = read_uploaded_file_for(ALICE, str(docx_p))
    assert "# Kenya scaling plan" in out and "We aim to scale three innovations." in out
    assert "| Drought-tolerant maize | 8 |" in out

    pdf_p = _upload(ws, ALICE, "brief.pdf", _pdf_bytes(["Page one text", "Second page"]))
    out = read_uploaded_file_for(ALICE, str(pdf_p))
    assert "PDF with 2 page(s)." in out and "Page one text" in out and "--- Page 2 ---" in out

    txt_p = _upload(ws, ALICE, "notes.md", "# Notes\nhello")
    assert "hello" in read_uploaded_file_for(ALICE, str(txt_p))


def test_bare_file_name_resolves_in_the_callers_own_folder(ws):
    from synapsis.tools.read_uploaded_file import read_uploaded_file_for

    _upload(ws, ALICE, "data.csv", "a,b\n1,2\n")
    assert "| a | b |" in read_uploaded_file_for(ALICE, "data.csv")


# ---------------------------------------------------------------------------
# Other users' uploads and everything else are denied
# ---------------------------------------------------------------------------

def test_other_users_upload_is_denied(ws):
    from synapsis.tools.read_uploaded_file import UploadAccessError, read_uploaded_file_for

    bobs = _upload(ws, BOB, "secret.xlsx", _xlsx_bytes({"S": [["x"], [1]]}))
    with pytest.raises(UploadAccessError):
        read_uploaded_file_for(ALICE, str(bobs))
    # Traversal from Alice's folder into Bob's is denied as well.
    from synapsis.user_files import owner_area, owner_key
    sneaky = str(owner_area("uploads", ALICE, ws) / ".." / owner_key(BOB) / "secret.xlsx")
    with pytest.raises(UploadAccessError):
        read_uploaded_file_for(ALICE, sneaky)
    with pytest.raises(UploadAccessError):
        read_uploaded_file_for(ALICE, f"../{owner_key(BOB)}/secret.xlsx")
    # Bob himself can read it.
    assert "| x |" in read_uploaded_file_for(BOB, str(bobs))


def test_outputs_references_hidden_and_symlink_escapes_are_denied(ws, tmp_path):
    from synapsis.tools.read_uploaded_file import UploadAccessError, read_uploaded_file_for
    from synapsis.user_files import new_output_path, owner_area

    out = new_output_path(ALICE, "report", "md", ws)
    out.write_text("mine but not an upload")
    secret = tmp_path / "server-secret.txt"
    secret.write_text("TOP SECRET")
    _upload(ws, ALICE, "ok.txt", "fine")
    link = owner_area("uploads", ALICE, ws) / "link.txt"
    os.symlink(secret, link)
    hidden = _upload(ws, ALICE, ".env.txt", "KEY=1")
    for raw in (str(out), str(secret), str(link), str(hidden), "/etc/passwd", "", None):
        with pytest.raises(UploadAccessError):
            read_uploaded_file_for(ALICE, raw)


def test_tool_uses_the_connection_identity_not_tool_input(ws):
    from synapsis.auth.context import set_current_user_id
    from synapsis.tools.read_uploaded_file import read_uploaded_file

    bobs = _upload(ws, BOB, "b.csv", "k,v\nsecret,42\n")

    async def call(user, path):
        set_current_user_id(user, "researcher")
        return await read_uploaded_file.handler({"path": path})

    denied = asyncio.run(call(ALICE, str(bobs)))
    assert denied.get("is_error") is True and "not one of this user's own uploads" in denied["content"][0]["text"]
    assert "secret" not in denied["content"][0]["text"]
    allowed = asyncio.run(call(BOB, str(bobs)))
    assert not allowed.get("is_error") and "| secret | 42 |" in allowed["content"][0]["text"]


# ---------------------------------------------------------------------------
# Caps
# ---------------------------------------------------------------------------

def test_row_and_column_caps_with_truncation_note(ws):
    r = _module()

    header = [f"c{i}" for i in range(r.MAX_COLUMNS + 10)]
    rows = [header] + [[f"v{n}"] + [n] * (len(header) - 1) for n in range(r.MAX_ROWS + 50)]
    p = _upload(ws, ALICE, "big.xlsx", _xlsx_bytes({"Big": rows}))
    text = r.read_uploaded_file_for(ALICE, str(p))
    assert f"{r.MAX_ROWS + 50} data row(s) + 1 header row, {r.MAX_COLUMNS + 10} column(s)." in text
    assert f"only the first {r.MAX_ROWS} of {r.MAX_ROWS + 50} data rows are shown" in text
    assert f"only the first {r.MAX_COLUMNS} of {r.MAX_COLUMNS + 10} columns are shown" in text
    table_lines = [ln for ln in text.splitlines() if ln.startswith("| ")]
    assert len(table_lines) == 2 + r.MAX_ROWS  # header + separator + capped rows
    assert table_lines[0].count("|") == r.MAX_COLUMNS + 1
    assert f"v{r.MAX_ROWS - 1}" in text and f"v{r.MAX_ROWS}" not in text


def test_text_cap_and_file_size_cap(ws, monkeypatch):
    r = _module()

    p = _upload(ws, ALICE, "long.txt", "x" * (r.MAX_CHARS + 5000))
    text = r.read_uploaded_file_for(ALICE, str(p))
    assert f"showing the first {r.MAX_CHARS:,}" in text
    assert text.count("x") <= r.MAX_CHARS + 10
    monkeypatch.setattr(r, "MAX_FILE_BYTES", 100)
    big = _upload(ws, ALICE, "big.csv", "a,b\n" + "1,2\n" * 100)
    with pytest.raises(r.UploadReadError, match="too large"):
        r.read_uploaded_file_for(ALICE, str(big))


def test_unsupported_and_corrupt_files_give_clear_errors(ws):
    from synapsis.tools.read_uploaded_file import UploadReadError, read_uploaded_file_for

    with pytest.raises(UploadReadError, match="xlsx"):
        read_uploaded_file_for(ALICE, str(_upload(ws, ALICE, "old.xls", b"\xd0\xcf\x11\xe0")))
    with pytest.raises(UploadReadError, match="not supported"):
        read_uploaded_file_for(ALICE, str(_upload(ws, ALICE, "run.exe", b"MZ")))
    with pytest.raises(UploadReadError, match="not a valid Office document"):
        read_uploaded_file_for(ALICE, str(_upload(ws, ALICE, "fake.xlsx", b"not a zip")))
    with pytest.raises(UploadReadError, match="no such uploaded file"):
        read_uploaded_file_for(ALICE, "missing.csv")


# ---------------------------------------------------------------------------
# Registration
# ---------------------------------------------------------------------------

def test_tool_is_registered_allowed_and_documented():
    import synapsis.agent_options as ao
    from synapsis.agents.definitions import _STANDARD_TOOLS
    from synapsis.system_prompt import build_system_prompt
    from synapsis.tools import IA_TOOLS

    assert "mcp__synapsis__read_uploaded_file" in ao.ALLOWED_TOOLS
    assert "mcp__synapsis__read_uploaded_file" in _STANDARD_TOOLS
    assert any(getattr(t, "name", "") == "read_uploaded_file" for t in IA_TOOLS)
    assert "mcp__synapsis__read_uploaded_file" in build_system_prompt({})
